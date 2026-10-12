"""Server-side report generation for Owner > Stock & Replenishment —
PDF/CSV/XLSX, built fresh from the database on every request.

Two report builders, each reusing the EXACT business logic the console's
own tabs already read:
  - build_purchasing_rows() calls wt/depletion.py's ingredient_forecast(),
    the same function GET /console/forecasts/ingredients calls for the
    Needs Purchasing tab -- so this report can never recommend something
    different from what that tab shows for the same ingredient.
  - build_expiry_rows() runs the same fefo_batch query and
    annotate_batch_expiry() call GET /console/expiry runs for the Expiring
    Soon tab, with a couple of extra SELECTed columns (category, supplier,
    unit_cost) that route doesn't need -- the batch selection and day-count
    logic is identical, not re-derived.

Three plain writers below the two builders turn either report's rows into
bytes: csv (stdlib csv.writer, safely quoted), xlsx (openpyxl, real numeric
types), pdf (reportlab, paginated automatically by Platypus' Table/
SimpleDocTemplate). Nothing in this module invents a number: a missing unit
cost stays None all the way through and prints as "not available", never a
fabricated 0 or an average over less data than it claims.

Every export is generated here, from a fresh query, at request time. The
frontend never supplies row data -- only a scope/filter/format selection.
"""
from __future__ import annotations

import csv
import io
import sqlite3
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    Paragraph, SimpleDocTemplate, Spacer, Table as RLTable, TableStyle,
)

from . import timeutil, unli
from .depletion import annotate_batch_expiry, ingredient_forecast

NO_SUPPLIER = "No supplier assigned"
PESO = "₱"


def _demand_model(conn: sqlite3.Connection) -> str:
    """Mirrors wt/api.py's own _demand_model() exactly (same app_setting,
    same thresholds) -- duplicated rather than imported because that
    helper is a private module-level function of api.py, not part of
    depletion.py's public surface. Both read the identical
    'depletion_demand_model' app_setting, so they can never disagree."""
    value = unli.get_setting(conn, "depletion_demand_model", 2.0)
    if value >= 2.0:
        return "hybrid"
    return "prophet" if value >= 1.0 else "ewma"


def _money(v: float | None) -> str:
    return f"{PESO}{v:,.2f}" if v is not None else "not available"


# ===========================================================================
# Purchasing-recommendation report
# ===========================================================================

def _needs_buying(r: dict) -> bool:
    # Same rule WhatToBuy.js's needsBuying() uses on the frontend, against
    # the same status/belowReorderPoint fields ingredient_forecast() already
    # computed -- not a second threshold invented here.
    return r["status"] in ("out", "critical", "at_risk") or bool(r["belowReorderPoint"])


def build_purchasing_rows(
    conn: sqlite3.Connection,
    scope: str = "needs_purchasing",
    ingredient_ids: list[int] | None = None,
) -> dict:
    """Rows + summary for the purchasing-recommendation report.

    scope:
      needs_purchasing (default) -- status out/critical/at_risk or below
                                     the configured reorder point
      urgent_only                -- status out/critical only
      selected                   -- exactly the given ingredient_ids
      all                        -- every deductable ingredient
    """
    forecast = ingredient_forecast(conn, demand_model=_demand_model(conn))
    ingredients = forecast["ingredients"]

    if scope == "urgent_only":
        chosen = [r for r in ingredients if r["status"] in ("out", "critical")]
    elif scope == "selected":
        wanted = set(ingredient_ids or [])
        chosen = [r for r in ingredients if r["ingredientId"] in wanted]
    elif scope == "all":
        chosen = list(ingredients)
    else:
        scope = "needs_purchasing"
        chosen = [r for r in ingredients if _needs_buying(r)]

    # ingredient_forecast()'s own row never carries unit_cost (it isn't part
    # of the depletion chain) -- a direct, fresh read of the ingredient
    # table is the only honest source for it here.
    cost_by_id: dict[int, float | None] = {}
    if chosen:
        ids = [r["ingredientId"] for r in chosen]
        qmarks = ",".join("?" * len(ids))
        for row in conn.execute(
            f"SELECT id, unit_cost FROM ingredient WHERE id IN ({qmarks})", ids,
        ):
            cost_by_id[row["id"]] = row["unit_cost"]

    rows = []
    for r in chosen:
        unit_cost = cost_by_id.get(r["ingredientId"])
        order_qty = r["suggestedOrderQty"]
        est_cost = order_qty * unit_cost if (order_qty is not None and unit_cost is not None) else None
        rows.append({
            "ingredient_id": r["ingredientId"],
            "ingredient": r["name"],
            "category": r.get("category") or None,
            "supplier": r.get("supplier") or NO_SUPPLIER,
            "current_usable_stock": r["available"],
            "base_unit": r["baseUnit"],
            "expected_daily_consumption": r["forecastDailyConsumption"] or None,
            "stock_cover_days": r["daysOfSupply"],
            "lead_time_days": r["leadTimeDays"],
            "safety_stock": r["safetyStockBase"],
            "suggested_purchase_qty": order_qty,
            "purchase_unit": r.get("suggestedOrderUnit") or r.get("purchaseUnit"),
            "unit_cost": unit_cost,
            "estimated_purchase_cost": est_cost,
            "priority": r["status"],
            "reason": r.get("replenishmentMessage") or "",
        })

    urgent_count = sum(1 for r in ingredients if r["status"] in ("out", "critical"))
    needs_count = sum(1 for r in ingredients if _needs_buying(r))
    known_costs = [r["estimated_purchase_cost"] for r in rows if r["estimated_purchase_cost"] is not None]
    missing_cost_rows = sum(1 for r in rows if r["estimated_purchase_cost"] is None)
    total_cost = sum(known_costs) if known_costs else None

    return {
        "rows": rows,
        "scope": scope,
        "summary": {
            "ingredients_to_purchase": needs_count,
            "urgent_items": urgent_count,
            "rows_in_report": len(rows),
            "estimated_total_purchase_cost": total_cost,
            # True whenever the total above is a sum over LESS than every
            # row in the report -- the caveat the PDF/summary must show
            # rather than presenting a partial sum as the whole picture.
            "cost_is_partial": missing_cost_rows > 0 and total_cost is not None,
            "cost_unavailable": total_cost is None and len(rows) > 0,
        },
    }


# ===========================================================================
# Expiry report
# ===========================================================================

_STATUS_LABEL = {"Expired": "Expired", "Use Soon": "Expiring soon", "Normal": "OK"}


def build_expiry_rows(
    conn: sqlite3.Connection,
    filter: str = "within_7",
    date_from: str | None = None,
    date_to: str | None = None,
    include_missing_expiry: bool = False,
) -> dict:
    """Rows + summary for the expiry report. Same open-batch population
    GET /console/expiry reads (fefo_batch state='open', qty_base>0) and the
    same annotate_batch_expiry() day-count/status logic -- category,
    supplier and unit_cost are extra SELECTed columns that route doesn't
    need, nothing about the expiry calculation itself is re-derived.

    filter: expired_only | within_3 | within_7 | within_14 | within_30 | custom
    (custom uses date_from/date_to, both Manila ISO dates, inclusive).
    """
    raw = conn.execute(
        """
        SELECT b.id, b.lot_no, b.qty_base, b.received_at, b.expires_at,
               i.id AS ingredient_id, i.name AS ingredient_name, i.base_unit,
               i.category, i.supplier, i.unit_cost,
               i.expiry_mode, i.shelf_life_days
          FROM fefo_batch b JOIN ingredient i ON i.id = b.ingredient_id
         WHERE b.state = 'open' AND b.qty_base > 0
        """
    ).fetchall()

    today = timeutil.business_today()
    rows = []
    for raw_row in raw:
        b = dict(raw_row)
        has_signal = b["expires_at"] is not None or b["expiry_mode"] == "shelf_life_estimate"
        if not has_signal:
            if not include_missing_expiry:
                continue
            b.pop("expiry_mode", None)
            b.pop("shelf_life_days", None)
            tracking, days_to_expiry, status = "none", None, None
        else:
            annotated = annotate_batch_expiry(dict(b))
            if annotated is None:
                if not include_missing_expiry:
                    continue
                b.pop("expiry_mode", None)
                b.pop("shelf_life_days", None)
                tracking, days_to_expiry, status = "none", None, None
            else:
                b = annotated
                tracking = b["tracking"]
                days_to_expiry = b["days_to_expiry"]
                status = b["storage_status"]

        if not _passes_filter(filter, days_to_expiry, b.get("expires_at"), date_from, date_to, today):
            continue

        unit_cost = b.get("unit_cost")
        qty = b.get("qty_base")
        value_at_risk = qty * unit_cost if (qty is not None and unit_cost is not None) else None
        rows.append({
            "ingredient": b["ingredient_name"],
            "category": b.get("category"),
            "batch_lot": b.get("lot_no") or f"#{b['id']}",
            "remaining_qty": qty,
            "base_unit": b["base_unit"],
            "expiry_date": b.get("expires_at"),
            "days_to_expiry": days_to_expiry,
            "status": "Expiry not recorded" if tracking == "none" else _STATUS_LABEL.get(status, status),
            "value_at_risk": value_at_risk,
            "supplier": b.get("supplier") or NO_SUPPLIER,
        })

    # Expired first, then soonest -- same ordering the Expiring Soon tab uses.
    rows.sort(key=lambda r: (r["days_to_expiry"] is None, r["days_to_expiry"] if r["days_to_expiry"] is not None else 0))

    expired_count = sum(1 for r in rows if r["days_to_expiry"] is not None and r["days_to_expiry"] < 0)
    soon_count = sum(1 for r in rows if r["days_to_expiry"] is not None and 0 <= r["days_to_expiry"] <= 7)
    missing_count = sum(1 for r in rows if r["days_to_expiry"] is None)
    known_values = [r["value_at_risk"] for r in rows if r["value_at_risk"] is not None]
    missing_value_rows = sum(1 for r in rows if r["value_at_risk"] is None)
    total_value = sum(known_values) if known_values else None

    return {
        "rows": rows,
        "filter": filter,
        "summary": {
            "batches_in_report": len(rows),
            "expired": expired_count,
            "expiring_soon": soon_count,
            "missing_expiry": missing_count,
            "estimated_total_value_at_risk": total_value,
            "value_is_partial": missing_value_rows > 0 and total_value is not None,
            "value_unavailable": total_value is None and len(rows) > 0,
        },
    }


def _passes_filter(filter: str, days: int | None, expires_at: str | None,
                    date_from: str | None, date_to: str | None, today) -> bool:
    if days is None:
        # Only reachable when include_missing_expiry=True and no computable
        # date exists at all -- always included (that's the point of the
        # flag), regardless of which date-bounded filter is selected.
        return True
    if filter == "expired_only":
        return days < 0
    if filter == "within_3":
        return days <= 3
    if filter == "within_7":
        return days <= 7
    if filter == "within_14":
        return days <= 14
    if filter == "within_30":
        return days <= 30
    if filter == "custom":
        if not expires_at:
            return False
        try:
            from datetime import date as _date
            d = _date.fromisoformat(str(expires_at)[:10])
        except ValueError:
            return False
        if date_from and d < _date.fromisoformat(date_from):
            return False
        if date_to and d > _date.fromisoformat(date_to):
            return False
        return True
    return True


# ===========================================================================
# Writers — csv / xlsx / pdf
# ===========================================================================

def write_csv(headers: list[str], rows: list[list[Any]]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, quoting=csv.QUOTE_MINIMAL)
    w.writerow(headers)
    for row in rows:
        w.writerow(row)
    return buf.getvalue().encode("utf-8-sig")  # BOM so Excel renders ₱ correctly


def write_xlsx(headers: list[str], rows: list[list[Any]], numeric_cols: set[int], sheet_title: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31] or "Report"

    bold = Font(bold=True)
    for col_idx, label in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.font = bold
        cell.alignment = Alignment(horizontal="left")

    for r_idx, row in enumerate(rows, start=2):
        for c_idx, value in enumerate(row, start=1):
            cell = ws.cell(row=r_idx, column=c_idx)
            if (c_idx - 1) in numeric_cols and isinstance(value, (int, float)):
                cell.value = value
                cell.number_format = "#,##0.00"
            elif value is None:
                cell.value = "not available"
            else:
                cell.value = value

    for col_idx, label in enumerate(headers, start=1):
        letter_ = get_column_letter(col_idx)
        longest = max([len(str(label))] + [len(str(row[col_idx - 1])) for row in rows] or [0])
        ws.column_dimensions[letter_].width = min(max(longest + 2, 10), 42)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


_PDF_STYLES = getSampleStyleSheet()
_TITLE_STYLE = ParagraphStyle("WTTitle", parent=_PDF_STYLES["Title"], fontSize=16, spaceAfter=2)
_SUB_STYLE = ParagraphStyle("WTSub", parent=_PDF_STYLES["Normal"], fontSize=9, textColor=colors.grey)
_H2_STYLE = ParagraphStyle("WTH2", parent=_PDF_STYLES["Heading2"], fontSize=11, spaceBefore=10, spaceAfter=4)
_BODY_STYLE = ParagraphStyle("WTBody", parent=_PDF_STYLES["Normal"], fontSize=9)
_NOTICE_STYLE = ParagraphStyle("WTNotice", parent=_PDF_STYLES["Normal"], fontSize=9,
                                textColor=colors.HexColor("#8a1f11"), spaceBefore=8)


def write_pdf(
    report_name: str,
    scope_line: str,
    summary_lines: list[str],
    headers: list[str],
    groups: list[tuple[str | None, list[list[Any]]]],
    notice: str | None = None,
) -> bytes:
    """groups: [(group_label_or_None, rows), ...] -- one table per group
    when grouping by supplier, or a single (None, rows) entry otherwise.
    Table pagination across pages is handled automatically by Platypus
    (SimpleDocTemplate splits a long Table flowable across page breaks)."""
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=landscape(letter),
        topMargin=0.5 * inch, bottomMargin=0.5 * inch,
        leftMargin=0.5 * inch, rightMargin=0.5 * inch,
    )
    now_manila = timeutil.to_manila(timeutil.utc_now())
    story: list[Any] = [
        Paragraph("WINGS THEORY", _TITLE_STYLE),
        Paragraph(report_name, _H2_STYLE),
        Paragraph(f"Generated: {now_manila.strftime('%Y-%m-%d %H:%M')} (Asia/Manila)", _SUB_STYLE),
        Paragraph("Prepared for: Owner", _SUB_STYLE),
        Paragraph(scope_line, _SUB_STYLE),
        Spacer(1, 10),
    ]
    for line in summary_lines:
        story.append(Paragraph(line, _BODY_STYLE))
    story.append(Spacer(1, 10))

    if notice:
        story.append(Paragraph(notice, _NOTICE_STYLE))
        story.append(Spacer(1, 10))

    any_rows = any(rows for _, rows in groups)
    if not any_rows:
        story.append(Paragraph("No matching records for this report.", _BODY_STYLE))
    else:
        for i, (label, rows) in enumerate(groups):
            if not rows:
                continue
            if label:
                story.append(Paragraph(label, _H2_STYLE))
            table_data = [headers] + [[_cell(v) for v in row] for row in rows]
            tbl = RLTable(table_data, repeatRows=1)
            tbl.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f3ead9")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 7.5),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#d8cdb8")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#fbf8f2")]),
            ]))
            story.append(tbl)
            if i < len(groups) - 1:
                story.append(Spacer(1, 14))

    doc.build(story)
    return buf.getvalue()


def _cell(v: Any) -> str:
    if v is None:
        return "not available"
    return str(v)


# ===========================================================================
# Report assembly — ties a builder's rows to a format writer
# ===========================================================================

PURCHASING_HEADERS = [
    "Ingredient", "Category", "Supplier", "Current Usable Stock", "Unit",
    "Expected Daily Consumption", "Stock Cover (days)", "Lead Time (days)",
    "Safety Stock", "Suggested Purchase Qty", "Purchase Unit",
    "Estimated Unit Cost", "Estimated Purchase Cost", "Priority", "Recommendation Reason",
]
PURCHASING_NUMERIC_COLS = {3, 5, 6, 7, 8, 9, 11, 12}  # 0-indexed into the row above


def _purchasing_row_values(r: dict, include_costs: bool) -> list[Any]:
    return [
        r["ingredient"], r["category"] or "not recorded", r["supplier"],
        r["current_usable_stock"], r["base_unit"],
        r["expected_daily_consumption"], r["stock_cover_days"], r["lead_time_days"],
        r["safety_stock"], r["suggested_purchase_qty"], r["purchase_unit"] or r["base_unit"],
        (r["unit_cost"] if include_costs else None),
        (r["estimated_purchase_cost"] if include_costs else None),
        r["priority"], r["reason"],
    ]


def generate_purchasing_report(
    conn: sqlite3.Connection, *, scope: str, ingredient_ids: list[int] | None,
    group_by_supplier: bool, fmt: str, include_costs: bool,
) -> tuple[bytes, str, str]:
    """Returns (file_bytes, filename, content_type)."""
    built = build_purchasing_rows(conn, scope, ingredient_ids)
    rows = built["rows"]
    summary = built["summary"]

    scope_label = {
        "needs_purchasing": "Needs Purchasing (everything currently recommended)",
        "urgent_only": "Urgent items only (out of stock / buy now)",
        "selected": "Selected ingredients",
        "all": "All ingredients",
    }.get(built["scope"], built["scope"])

    if fmt == "csv":
        values = [_purchasing_row_values(r, include_costs) for r in rows]
        content = write_csv(PURCHASING_HEADERS, values)
        return content, _filename("purchasing-list", "csv"), "text/csv; charset=utf-8"

    if fmt == "xlsx":
        values = [_purchasing_row_values(r, include_costs) for r in rows]
        content = write_xlsx(PURCHASING_HEADERS, values, PURCHASING_NUMERIC_COLS, "Purchasing List")
        return content, _filename("purchasing-list", "xlsx"), (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

    # pdf
    if group_by_supplier:
        by_supplier: dict[str, list[dict]] = {}
        for r in rows:
            by_supplier.setdefault(r["supplier"], []).append(r)
        groups = [
            (f"Supplier: {supplier} ({len(items)} item{'s' if len(items) != 1 else ''})",
             [_purchasing_row_values(r, include_costs) for r in items])
            for supplier, items in sorted(by_supplier.items(), key=lambda kv: kv[0])
        ]
    else:
        groups = [(None, [_purchasing_row_values(r, include_costs) for r in rows])]

    cost_line = "Estimated total purchase cost: not available (no cost data on the recommended items)"
    if summary["estimated_total_purchase_cost"] is not None:
        cost_line = f"Estimated total purchase cost: {_money(summary['estimated_total_purchase_cost'])}"
        if summary["cost_is_partial"]:
            cost_line += " (partial — some items have no recorded unit cost)"

    summary_lines = [
        f"Scope: {scope_label}",
        f"Ingredients to purchase: {summary['ingredients_to_purchase']}",
        f"Urgent items: {summary['urgent_items']}",
        f"Rows in this report: {summary['rows_in_report']}",
    ]
    if include_costs:
        summary_lines.append(cost_line)

    content = write_pdf(
        report_name="Purchasing Recommendation Report",
        scope_line=f"Scope: {scope_label}" + (" · Grouped by supplier" if group_by_supplier else ""),
        summary_lines=summary_lines,
        headers=PURCHASING_HEADERS if include_costs else PURCHASING_HEADERS[:11] + PURCHASING_HEADERS[13:],
        groups=[(label, [row if include_costs else row[:11] + row[13:] for row in group_rows])
                for label, group_rows in groups],
    )
    return content, _filename("purchasing-list", "pdf"), "application/pdf"


EXPIRY_HEADERS = [
    "Ingredient", "Category", "Batch/Lot", "Remaining Qty", "Base Unit",
    "Expiry Date", "Days Until Expiry", "Status", "Estimated Value at Risk", "Supplier",
]
EXPIRY_NUMERIC_COLS = {3, 6, 8}

EXPIRY_NOTICE = ("This report does not automatically discard expired stock. Actual stock "
                  "removal requires an authorized Admin action.")


def _expiry_row_values(r: dict) -> list[Any]:
    return [
        r["ingredient"], r["category"] or "not recorded", r["batch_lot"],
        r["remaining_qty"], r["base_unit"], r["expiry_date"] or "not recorded",
        r["days_to_expiry"], r["status"], r["value_at_risk"], r["supplier"],
    ]


def generate_expiry_report(
    conn: sqlite3.Connection, *, filter: str, date_from: str | None, date_to: str | None,
    include_missing_expiry: bool, fmt: str,
) -> tuple[bytes, str, str]:
    built = build_expiry_rows(conn, filter, date_from, date_to, include_missing_expiry)
    rows = built["rows"]
    summary = built["summary"]
    values = [_expiry_row_values(r) for r in rows]

    filter_label = {
        "expired_only": "Expired only",
        "within_3": "Expiring within 3 days",
        "within_7": "Expiring within 7 days",
        "within_14": "Expiring within 14 days",
        "within_30": "Expiring within 30 days",
        "custom": f"Custom range ({date_from or '...'} to {date_to or '...'})",
    }.get(filter, filter)

    if fmt == "csv":
        content = write_csv(EXPIRY_HEADERS, values)
        return content, _filename("expiry-report", "csv"), "text/csv; charset=utf-8"

    if fmt == "xlsx":
        content = write_xlsx(EXPIRY_HEADERS, values, EXPIRY_NUMERIC_COLS, "Expiry Report")
        return content, _filename("expiry-report", "xlsx"), (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

    value_line = "Estimated total value at risk: not available (no cost data on these batches)"
    if summary["estimated_total_value_at_risk"] is not None:
        value_line = f"Estimated total value at risk: {_money(summary['estimated_total_value_at_risk'])}"
        if summary["value_is_partial"]:
            value_line += " (partial — some batches have no recorded unit cost)"

    summary_lines = [
        f"Filter: {filter_label}" + (" · including batches with no recorded expiry" if include_missing_expiry else ""),
        f"Batches in this report: {summary['batches_in_report']}",
        f"Expired: {summary['expired']}",
        f"Expiring within 7 days: {summary['expiring_soon']}",
        f"Missing expiry data: {summary['missing_expiry']}",
        value_line,
    ]

    content = write_pdf(
        report_name="Expiry Report",
        scope_line=f"Filter: {filter_label}",
        summary_lines=summary_lines,
        headers=EXPIRY_HEADERS,
        groups=[(None, values)],
        notice=EXPIRY_NOTICE,
    )
    return content, _filename("expiry-report", "pdf"), "application/pdf"


def _filename(base: str, ext: str) -> str:
    today = timeutil.to_manila(timeutil.utc_now()).strftime("%Y-%m-%d")
    return f"wings-theory-{base}-{today}.{ext}"
